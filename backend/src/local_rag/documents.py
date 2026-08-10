from __future__ import annotations

import re
from pathlib import Path

from .models import Chunk

SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".xlsx"}


def discover_documents(path: Path) -> list[Path]:
    if path.is_file():
        candidates = [path]
    elif path.is_dir():
        candidates = [item for item in path.rglob("*") if item.is_file()]
    else:
        raise FileNotFoundError(f"Belge yolu bulunamadı: {path}")
    return sorted(item for item in candidates if item.suffix.lower() in SUPPORTED_SUFFIXES)


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF okumak için pypdf kurulmalıdır.") from exc
        pages = []
        for number, page in enumerate(PdfReader(str(path)).pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"[Sayfa {number}]\n{text}")
        return "\n\n".join(pages)
    if suffix == ".xlsx":
        return _read_excel_workbook(path)
    raise ValueError(f"Desteklenmeyen belge türü: {path.suffix}")


def _read_excel_workbook(path: Path) -> str:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Excel okumak için openpyxl kurulmalıdır.") from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Excel dosyası okunamadı: {path.name}") from exc

    sections: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row_number, cells in enumerate(worksheet.iter_rows(), start=1):
                values: list[str] = []
                for column_number, cell in enumerate(cells, start=1):
                    value = cell.value
                    if value is None or str(value).strip() == "":
                        continue
                    coordinate = f"{get_column_letter(column_number)}{row_number}"
                    values.append(f"{coordinate}={_format_excel_value(value)}")
                if values:
                    rows.append(" | ".join(values))
            if rows:
                sections.append(
                    f"[Çalışma Sayfası: {worksheet.title}]\n" + "\n".join(rows)
                )
    finally:
        workbook.close()
    return "\n\n".join(sections)


def _format_excel_value(value: object) -> str:
    if hasattr(value, "isoformat"):
        try:
            return str(value.isoformat())
        except (TypeError, ValueError):
            pass
    if isinstance(value, float):
        return f"{value:.10g}"
    return re.sub(r"\s+", " ", str(value)).strip()


def chunk_text(text: str, source: str, chunk_size: int = 1200, overlap: int = 200) -> list[Chunk]:
    if chunk_size < 100:
        raise ValueError("chunk_size en az 100 olmalıdır")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap, 0 ile chunk_size arasında olmalıdır")
    normalized = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n"))
    normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()
    if not normalized:
        return []

    chunks: list[Chunk] = []
    start = 0
    position = 0
    while start < len(normalized):
        hard_end = min(start + chunk_size, len(normalized))
        end = hard_end
        if hard_end < len(normalized):
            search_from = start + chunk_size // 2
            boundaries = [normalized.rfind("\n\n", search_from, hard_end), normalized.rfind(". ", search_from, hard_end)]
            boundary = max(boundaries)
            if boundary > start:
                end = boundary + (2 if normalized[boundary : boundary + 2] in {"\n\n", ". "} else 0)
        content = normalized[start:end].strip()
        if content:
            chunks.append(Chunk(source=source, position=position, content=content))
            position += 1
        if end >= len(normalized):
            break
        start = max(end - overlap, start + 1)
    return chunks
